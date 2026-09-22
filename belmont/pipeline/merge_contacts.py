#!/usr/bin/env python3
"""Append phone numbers and emails to a campaign from (a) skip-trace vendor result files and
(b) a hand-maintained manual_contacts.csv, then rebuild the campaign files and the dashboard.

Vendor files: drop every CSV a skip-tracing service returns (BatchSkipTracing, PropStream, REISkip,
DirectSkip, TLO-based vendors...) into  campaigns/<campaign>/skip_trace_results/ . Columns are detected
by name: anything containing phone / mobile / cell / landline is a phone, anything containing email is
an email, "dnc" or "do not call" columns flag phones. Rows are matched back to households by the
Contact ID column you uploaded (hh_...), or by last name + mailing address + ZIP.

Manual file: campaigns/<campaign>/manual_contacts.csv (contact_id, full_name, phone, email, source,
do_not_call, notes) for contact details you collect yourself from replies and calls.

Nothing this script writes is committed: contacts_appended.csv and mailing_list_with_contacts.* are
ignored by git, and the dashboard receives contacts only through its private shared database
(paste contacts_appended.csv into Campaign > Sender & letter text > Import contacts).

Usage:  python merge_contacts.py --campaign life_estate
"""
import argparse
import csv
import glob
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, ".."))
CAMP = os.path.join(ROOT, "campaigns")


def norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def norm_phone(p):
    d = re.sub(r"\D", "", p or "")
    if len(d) == 11 and d.startswith("1"):
        d = d[1:]
    if len(d) != 10:
        return ""
    return f"({d[:3]}) {d[3:6]}-{d[6:]}"


def norm_email(e):
    e = (e or "").strip().lower()
    return e if re.fullmatch(r"[^@\s]+@[^@\s]+\.[a-z]{2,}", e) else ""


def load_households(camp):
    path = os.path.join(CAMP, camp, "mailing_list.csv")
    rows = list(csv.DictReader(open(path, newline="", encoding="utf-8")))
    by_id = {r["contact_id"]: r for r in rows}
    by_key = {}
    for r in rows:
        by_key[(norm(r["last_name"]), norm(r["address_line_1"]), (r["zip"] or "")[:5])] = r["contact_id"]
        by_key[("prop", norm(r["property_address"]))] = r["contact_id"]
    return by_id, by_key


def match_row(row, headers, by_id, by_key):
    for h in headers:
        if norm(h) in {"contactid", "contact_id", "custom", "custom1", "customfield", "reference", "ref", "id"} and str(row.get(h, "")).startswith("hh_"):
            return row[h].strip()
    def col(*names):
        for h in headers:
            if norm(h) in {norm(n) for n in names}:
                return row.get(h, "")
        return ""
    last = col("Last Name", "LastName", "Owner Last Name")
    addr = col("Mailing Address", "Mail Address", "Address", "Mailing Street")
    zip5 = (col("Mailing Zip", "Mail Zip", "Zip", "Mailing ZIP Code") or "")[:5]
    hit = by_key.get((norm(last), norm(addr), zip5))
    if not hit:
        hit = by_key.get(("prop", norm(col("Property Address", "Property Street", "Site Address"))))
    return hit


def merge(camp):
    by_id, by_key = load_households(camp)
    found = {cid: {"phones": [], "emails": [], "dnc": [], "sources": []} for cid in by_id}
    unmatched = 0
    for path in sorted(glob.glob(os.path.join(CAMP, camp, "skip_trace_results", "*.csv"))):
        src = os.path.basename(path)
        with open(path, newline="", encoding="utf-8-sig") as f:
            rd = csv.DictReader(f)
            headers = rd.fieldnames or []
            phone_cols = [h for h in headers if re.search(r"phone|mobile|cell|landline|\btel", h, re.I) and not re.search(r"type|carrier|dnc|score|status", h, re.I)]
            email_cols = [h for h in headers if re.search(r"e-?mail", h, re.I)]
            dnc_cols = [h for h in headers if re.search(r"dnc|do.?not.?call", h, re.I)]
            for row in rd:
                cid = match_row(row, headers, by_id, by_key)
                if not cid or cid not in found:
                    unmatched += 1
                    continue
                rec = found[cid]
                for i, h in enumerate(phone_cols):
                    p = norm_phone(row.get(h, ""))
                    if p and p not in rec["phones"]:
                        rec["phones"].append(p)
                    if p:
                        # per-phone DNC column ("Phone 1 DNC") or a row-level flag
                        flag = False
                        for d in dnc_cols:
                            m = re.search(r"(\d)", d)
                            if (m and int(m.group(1)) == i + 1) or not m:
                                if str(row.get(d, "")).strip().lower() in {"y", "yes", "true", "1", "dnc"}:
                                    flag = True
                        if flag and p not in rec["dnc"]:
                            rec["dnc"].append(p)
                for h in email_cols:
                    e = norm_email(row.get(h, ""))
                    if e and e not in rec["emails"]:
                        rec["emails"].append(e)
                if src not in rec["sources"]:
                    rec["sources"].append(src)
    manual = os.path.join(CAMP, camp, "manual_contacts.csv")
    if os.path.exists(manual):
        for row in csv.DictReader(open(manual, newline="", encoding="utf-8-sig")):
            cid = (row.get("contact_id") or "").strip()
            if cid not in found:
                continue
            rec = found[cid]
            p, e = norm_phone(row.get("phone", "")), norm_email(row.get("email", ""))
            if p and p not in rec["phones"]:
                rec["phones"].append(p)
            if p and str(row.get("do_not_call", "")).strip().lower() in {"y", "yes", "true", "1"} and p not in rec["dnc"]:
                rec["dnc"].append(p)
            if e and e not in rec["emails"]:
                rec["emails"].append(e)
            s = (row.get("source") or "manual").strip()
            if (p or e) and s not in rec["sources"]:
                rec["sources"].append(s)
    out = os.path.join(CAMP, camp, "contacts_appended.csv")
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["contact_id", "full_name", "phones", "emails", "do_not_call_phones", "sources"])
        n = 0
        for cid, rec in found.items():
            if rec["phones"] or rec["emails"]:
                n += 1
                w.writerow([cid, by_id[cid]["full_name"], "; ".join(rec["phones"]), "; ".join(rec["emails"]), "; ".join(rec["dnc"]), "; ".join(rec["sources"])])
    # private copies of the mailing list with contact columns (both ignored by git)
    ml = list(csv.DictReader(open(os.path.join(CAMP, camp, "mailing_list.csv"), newline="", encoding="utf-8")))
    cols = list(ml[0].keys()) + ["phone_1", "phone_2", "phone_3", "email_1", "email_2", "do_not_call_phones", "contact_sources"] if ml else []
    with open(os.path.join(CAMP, camp, "mailing_list_with_contacts.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in ml:
            rec = found.get(r["contact_id"], {"phones": [], "emails": [], "dnc": [], "sources": []})
            ph, em = rec["phones"] + ["", "", ""], rec["emails"] + ["", ""]
            r.update({"phone_1": ph[0], "phone_2": ph[1], "phone_3": ph[2], "email_1": em[0], "email_2": em[1],
                      "do_not_call_phones": "; ".join(rec["dnc"]), "contact_sources": "; ".join(rec["sources"])})
            w.writerow(r)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook(); ws = wb.active; ws.title = "Households + contacts"
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            ws.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF"); ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="1F3A5F")
        for r in csv.DictReader(open(os.path.join(CAMP, camp, "mailing_list_with_contacts.csv"), newline="", encoding="utf-8")):
            ws.append([r[c] for c in cols])
        ws.freeze_panes = "C2"
        wb.save(os.path.join(CAMP, camp, "mailing_list_with_contacts.xlsx"))
    except Exception as e:  # noqa: BLE001
        print("xlsx not written:", e)
    print(f"{camp}: contact info for {n} of {len(by_id)} households ({unmatched} vendor rows could not be matched)")
    print(f"  -> {out}\n  -> mailing_list_with_contacts.csv / .xlsx (private, ignored by git)")
    print("  Paste contacts_appended.csv into the dashboard (Campaign > Sender & letter text > Import contacts) to see phones and emails on the cards.")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--campaign", default="life_estate")
    a = ap.parse_args()
    merge(a.campaign)


if __name__ == "__main__":
    main()
